from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import BudgetSnapshot, ScenarioResult, Transaction

SAGE = "0F3D2E"
WHITE = "FFFFFF"
ROW_ALT = "EFE8DA"
RED = "9B2C2C"
GREEN = "1F6F4A"


def write_workbook(
    path: Path,
    transactions: list[Transaction],
    budget: BudgetSnapshot,
    scenario: ScenarioResult | None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _overview(wb.active, budget, scenario)
    _transactions(wb.create_sheet("Transactions"), transactions)
    _budget(wb.create_sheet("Budget"), budget)
    _scenario(wb.create_sheet("Scenario"), scenario)
    wb.save(path)


def _header_fill() -> PatternFill:
    return PatternFill("solid", fgColor=SAGE)


def _header_font() -> Font:
    return Font(name="Calibri", bold=True, color=WHITE, size=12)


def _money(cell, value: float) -> None:
    cell.value = round(value, 2)
    cell.number_format = '"$"#,##0.00'


def _style_header(ws, cols: int, row: int = 1) -> None:
    thin = Border(bottom=Side(style="thin", color="D6D0C4"))
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = _header_fill()
        cell.font = _header_font()
        cell.alignment = Alignment(horizontal="left")
        cell.border = thin


def _autosize(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _overview(ws, budget: BudgetSnapshot, scenario: ScenarioResult | None) -> None:
    ws.title = "Overview"
    ws["A1"] = "MultiPly budget overview"
    ws["A1"].font = Font(name="Calibri", bold=True, size=18, color=SAGE)
    ws["A2"] = "Educational simulation, not financial advice."
    ws["A2"].font = Font(italic=True, color="6B6458")

    labels = [
        ("Monthly take-home", budget.income_monthly),
        ("Monthly spending", budget.burn_monthly),
        ("Surplus before investing", budget.surplus_monthly),
        ("Planned investing", budget.invest_monthly),
        ("Cash on hand", budget.cash_on_hand),
    ]
    ws["A4"] = "Metric"
    ws["B4"] = "Amount"
    _style_header(ws, 2, row=4)
    for offset, (label, amount) in enumerate(labels, start=5):
        ws.cell(offset, 1, label)
        _money(ws.cell(offset, 2), amount)
    if scenario:
        ws["A11"] = "Latest scenario"
        ws["A11"].font = Font(bold=True, color=SAGE)
        ws["A12"] = scenario.question
        ws["A13"] = "Affordable"
        ws["B13"] = "Yes" if scenario.affordable else "Not yet"
        ws["B13"].font = Font(bold=True, color=GREEN if scenario.affordable else RED)
        ws["A14"] = "Rent cap"
        _money(ws.cell(14, 2), scenario.rent_cap)
        ws["A15"] = "Target rent"
        _money(ws.cell(15, 2), scenario.target_rent)
        ws["A16"] = "Runway (months)"
        ws["B16"] = scenario.runway_months
    _autosize(ws, [32, 18, 40])


def _transactions(ws, transactions: list[Transaction]) -> None:
    headers = ["txn_id", "date", "merchant", "amount", "category", "is_recurring"]
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header)
    _style_header(ws, len(headers))
    seen: set[str] = set()
    row = 2
    for txn in sorted(transactions, key=lambda item: item.date, reverse=True):
        if txn.txn_id in seen:
            continue
        seen.add(txn.txn_id)
        ws.cell(row, 1, txn.txn_id)
        ws.cell(row, 2, txn.date)
        ws.cell(row, 3, txn.merchant)
        _money(ws.cell(row, 4), txn.amount)
        ws.cell(row, 5, txn.category or "Uncategorized")
        ws.cell(row, 6, "yes" if txn.is_recurring else "no")
        if row % 2 == 0:
            for col in range(1, 7):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=ROW_ALT)
        row += 1
    _autosize(ws, [14, 14, 28, 14, 16, 14])


def _budget(ws, budget: BudgetSnapshot) -> None:
    ws["A1"] = "Category"
    ws["B1"] = "Monthly amount"
    _style_header(ws, 2)
    ws.cell(2, 1, "Take-home income")
    _money(ws.cell(2, 2), budget.income_monthly)
    row = 3
    for item in budget.categories:
        ws.cell(row, 1, item.name)
        _money(ws.cell(row, 2), item.amount)
        row += 1
    ws.cell(row, 1, "Total spending")
    ws.cell(row, 1).font = Font(bold=True)
    _money(ws.cell(row, 2), budget.burn_monthly)
    row += 1
    ws.cell(row, 1, "Surplus")
    ws.cell(row, 1).font = Font(bold=True)
    _money(ws.cell(row, 2), budget.surplus_monthly)
    _autosize(ws, [28, 18])


def _scenario(ws, scenario: ScenarioResult | None) -> None:
    ws["A1"] = "Field"
    ws["B1"] = "Value"
    _style_header(ws, 2)
    if scenario is None:
        ws["A2"] = "No scenario asked yet"
        return
    rows = [
        ("Question", scenario.question),
        ("Affordable", "Yes" if scenario.affordable else "Not yet"),
        ("Current rent", scenario.current_rent),
        ("Target rent", scenario.target_rent),
        ("Rent cap", scenario.rent_cap),
        ("Invest monthly", scenario.invest_monthly),
        ("Leftover monthly", scenario.leftover_monthly),
        ("Runway months", scenario.runway_months),
        ("Cash needed for 3-month runway", scenario.cash_needed_for_runway),
        ("Gap", scenario.gap),
    ]
    for index, (label, value) in enumerate(rows, start=2):
        ws.cell(index, 1, label)
        cell = ws.cell(index, 2, value if isinstance(value, str) else None)
        if isinstance(value, float):
            _money(cell, value)
        elif isinstance(value, str):
            cell.value = value
    ws.cell(len(rows) + 3, 1, "Assumptions")
    ws.cell(len(rows) + 3, 1).font = Font(bold=True, color=SAGE)
    for offset, assumption in enumerate(scenario.assumptions, start=len(rows) + 4):
        ws.cell(offset, 1, assumption)
    _autosize(ws, [36, 80])
